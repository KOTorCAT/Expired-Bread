"""Демонстрационные данные: два пользователя и несколько документов.

    python manage.py seed_demo
    python manage.py seed_demo --reset   # удалить прежние демо-документы

Документы создаются сразу в статусе ready с заполненными data и полями
индекса, поэтому поиск и фильтры работают без обращения к модели.
Файлы генерируются на месте в формате DOCX — их можно скачать и открыть.
"""

from __future__ import annotations

import io

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand
from docx import Document as DocxDocument

from documents.models import Document
from documents.services.field_index import index_document
from documents.services.normalization import build_search_text

DEMO_USERS = (
    ("demo", "demo12345"),
    ("other", "other12345"),
)

DEMO_DOCUMENTS = (
    {
        "owner": "demo",
        "doc_type": "price_list",
        "filename": "prajs-list-2025.docx",
        "text": (
            "ООО «ТехноСнаб»\nПрайс-лист от 15.03.2025\n\n"
            "Ноутбук ASUS VivoBook 15 | 54990 руб. | 15 шт\n"
            "Ноутбук Lenovo IdeaPad 3 | 47990 руб. | 8 шт\n"
            "Монитор Dell P2422H | 18990 руб. | 12 шт\n"
        ),
        "data": {
            "title": "Прайс-лист ООО «ТехноСнаб»",
            "description": "Ноутбуки и мониторы с ценами и остатками.",
            "organization": "ООО «ТехноСнаб»",
            "date": "15.03.2025",
            "vendor": "ТехноСнаб",
            "currency": "RUB",
            "items": [
                {"name": "Ноутбук ASUS VivoBook 15", "sku": "AS-15",
                 "price": "54 990 руб.", "unit": "шт"},
                {"name": "Ноутбук Lenovo IdeaPad 3", "sku": "LN-IP3",
                 "price": "47 990 руб.", "unit": "шт"},
                {"name": "Монитор Dell P2422H", "sku": "DL-P2422H",
                 "price": "18 990 руб.", "unit": "шт"},
            ],
        },
    },
    {
        "owner": "demo",
        "doc_type": "specification",
        "filename": "specification-noutbuk.docx",
        "text": (
            "Спецификация № 12 от 02.04.2025\n\n"
            "Изделие: ноутбук ASUS VivoBook 15\n"
            "Производитель: ASUS\n"
            "Оперативная память: 16 ГБ\n"
            "Накопитель: 512 ГБ SSD\n"
            "Цена: 54 990 руб.\n"
        ),
        "data": {
            "title": "Спецификация на ноутбук ASUS VivoBook 15",
            "description": "Характеристики и цена ноутбука.",
            "organization": "ООО «ТехноСнаб»",
            "date": "02.04.2025",
            "product_name": "Ноутбук ASUS VivoBook 15",
            "manufacturer": "ASUS",
            "category": "Ноутбуки",
            "price": "54 990 руб.",
            "ram": "16 ГБ",
            "storage": "512 ГБ SSD",
            "characteristics": {
                "Диагональ": "15.6 дюйма",
                "Процессор": "Intel Core i5-1235U",
            },
        },
    },
    {
        "owner": "other",
        "doc_type": "specification",
        "filename": "specification-monitor.docx",
        "text": (
            "Спецификация № 7\n\n"
            "Изделие: монитор Dell P2422H\n"
            "Производитель: Dell\n"
            "Цена: 18 990 руб.\n"
        ),
        "data": {
            "title": "Спецификация на монитор Dell P2422H",
            "description": "Характеристики и цена монитора.",
            "organization": "ООО «ТехноСнаб»",
            "date": "10.02.2025",
            "product_name": "Монитор Dell P2422H",
            "manufacturer": "Dell",
            "category": "Мониторы",
            "price": "18 990 руб.",
            "ram": "",
            "storage": "",
            "characteristics": {"Диагональ": "23.8 дюйма"},
        },
    },
)


def _docx_bytes(text: str) -> bytes:
    """Собирает простой DOCX в памяти."""
    document = DocxDocument()
    for line in text.splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class Command(BaseCommand):
    help = "Создаёт демо-пользователей и документы"

    def add_arguments(self, parser):
        parser.add_argument(
            "--reset",
            action="store_true",
            help="Удалить прежние демо-документы перед созданием",
        )

    def handle(self, *args, **options):
        User = get_user_model()
        users: dict[str, object] = {}

        for username, password in DEMO_USERS:
            user, created = User.objects.get_or_create(username=username)
            if created:
                user.set_password(password)
                user.save()
                self.stdout.write(f"Создан пользователь {username} / {password}")
            else:
                self.stdout.write(f"Пользователь {username} уже существует")
            users[username] = user

        if options["reset"]:
            removed, _ = Document.objects.filter(
                user__username__in=[name for name, _ in DEMO_USERS]
            ).delete()
            self.stdout.write(f"Удалено записей: {removed}")

        created_count = 0

        for spec in DEMO_DOCUMENTS:
            owner = users[spec["owner"]]

            exists = Document.objects.filter(
                user=owner,
                original_filename=spec["filename"],
            ).exists()
            if exists:
                self.stdout.write(
                    f"Документ {spec['filename']} уже есть — пропускаю"
                )
                continue

            document = Document(
                user=owner,
                original_filename=spec["filename"],
                document_type=spec["doc_type"],
                title=spec["data"]["title"],
                description=spec["data"].get("description", ""),
                data=spec["data"],
                data_raw=spec["data"],
                processing_meta={"demo": True, "warnings": []},
                extracted_text=spec["text"],
                status=Document.Status.READY,
                reviewed_at=None,
            )
            document.file.save(
                spec["filename"],
                ContentFile(_docx_bytes(spec["text"])),
                save=False,
            )
            document.search_text = build_search_text(
                title=document.title,
                description=document.description,
                original_filename=document.original_filename,
                extracted_text=document.extracted_text,
            )
            document.save()
            index_document(document)
            created_count += 1
            self.stdout.write(f"Создан документ: {spec['filename']}")

        self.stdout.write(
            self.style.SUCCESS(f"Готово. Добавлено документов: {created_count}")
        )
        self.stdout.write(
            "Вход: demo / demo12345 (и second пользователь other / other12345)"
        )
