"""Извлечение текста из PDF и DOCX.

Модуль не знает ничего о Django, моделях и пользователях.
Вход — путь к файлу, выход — строка с текстом.
"""

from __future__ import annotations

import os
from pathlib import Path

import pymupdf as fitz  # PyMuPDF 
from docx import Document as DocxDocument


class ExtractionError(Exception):
    """Не удалось извлечь текст из файла."""


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_text(file_path: str | os.PathLike, file_type: str | None = None) -> str:
    """Извлекает текст из файла.

    file_type можно не передавать — определим по расширению.
    """
    path = Path(file_path)

    if not path.exists():
        raise ExtractionError(f"Файл не найден: {path}")

    ext = (file_type or path.suffix).lower()
    if not ext.startswith("."):
        ext = "." + ext

    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(f"Неподдерживаемый формат: {ext}")

    if ext == ".pdf":
        return extract_pdf_text(path)
    return extract_docx_text(path)


def extract_pdf_text(file_path: str | os.PathLike) -> str:
    """Извлекает текст из PDF через PyMuPDF.

    Если PDF состоит из сканов, текстового слоя нет — вернётся пустая строка,
    и это ошибка, а не «успешно обработанный документ».
    """
    path = Path(file_path)
    parts: list[str] = []

    try:
        with fitz.open(path) as doc:
            for page in doc:
                parts.append(page.get_text())
    except Exception as exc:
        raise ExtractionError(f"Не удалось прочитать PDF: {exc}") from exc

    text = "\n".join(parts).strip()
    if not text:
        raise ExtractionError(
            "В PDF не найден текстовый слой. Возможно, это скан. "
            "OCR пока не поддерживается."
        )
    return text


def extract_docx_text(file_path: str | os.PathLike) -> str:
    """Извлекает текст из DOCX через python-docx."""
    path = Path(file_path)
    parts: list[str] = []

    try:
        doc = DocxDocument(str(path))
        for paragraph in doc.paragraphs:
            parts.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append(" | ".join(cells))
    except Exception as exc:
        raise ExtractionError(f"Не удалось прочитать DOCX: {exc}") from exc

    text = "\n".join(p for p in parts if p).strip()
    if not text:
        raise ExtractionError("DOCX не содержит текста.")
    return text